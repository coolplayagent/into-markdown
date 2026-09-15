#!/usr/bin/env python3
"""Build source-controlled cross-format OCR cases from the existing OCR golden.

Requires the document fixture libraries python-docx, python-pptx and openpyxl.
The recognized sentence occurs only in the PNG, never in container text or alt text.
Run this local qualification separately from the four fast CI jobs.
"""
import argparse
import base64
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import importlib.util
import zipfile

ROOT = Path(__file__).resolve().parents[2]
NATIVE = 'Native container text.'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def archive(path, entries):
    with zipfile.ZipFile(path, 'w') as output:
        for name, data in entries.items():
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_STORED if name == 'mimetype' else zipfile.ZIP_DEFLATED
            output.writestr(info, data)


def odf(path, image, family):
    media = 'application/vnd.oasis.opendocument.' + {'odt': 'text', 'ods': 'spreadsheet', 'odp': 'presentation'}[family]
    namespaces = ' '.join(f'xmlns:{name}="urn:oasis:names:tc:opendocument:xmlns:{value}:1.0"'
        for name, value in [('office', 'office'), ('text', 'text'), ('draw', 'drawing'),
                            ('table', 'table'), ('svg', 'svg-compatible')])
    namespaces += ' xmlns:xlink="http://www.w3.org/1999/xlink"'
    frame = '<draw:frame svg:x="0cm" svg:y="0cm" svg:width="20cm" svg:height="3cm"><draw:image xlink:type="simple" xlink:href="Pictures/scan.png"/></draw:frame>'
    body = f'<text:p>{NATIVE}</text:p>{frame}'
    if family == 'odt':
        body = f'<office:text>{body}</office:text>'
    elif family == 'odp':
        body = f'<office:presentation><draw:page draw:name="Scan">{frame}<draw:frame><draw:text-box><text:p>{NATIVE}</text:p></draw:text-box></draw:frame></draw:page></office:presentation>'
    else:
        body = f'<office:spreadsheet><table:table table:name="Scan"><table:table-row><table:table-cell office:value-type="string"><text:p>{NATIVE}</text:p>{frame}</table:table-cell></table:table-row></table:table></office:spreadsheet>'
    content = f'<office:document-content {namespaces} office:version="1.3"><office:body>{body}</office:body></office:document-content>'
    manifest = f'<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3"><manifest:file-entry manifest:full-path="/" manifest:media-type="{media}"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/><manifest:file-entry manifest:full-path="Pictures/scan.png" manifest:media-type="image/png"/></manifest:manifest>'
    archive(path, {'mimetype': media, 'content.xml': content, 'META-INF/manifest.xml': manifest,
                   'Pictures/scan.png': image})


def build(root, office=None):
    from docx import Document
    from docx.shared import Inches
    from pptx import Presentation
    from pptx.util import Inches as SlideInches
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image

    root.mkdir(parents=True, exist_ok=False)
    fixtures = json.loads((ROOT / 'fixtures/manifest.json').read_text())
    fixture = next(x for x in fixtures['fixtures'] if x['id'] == 'ocr-english-clear-1')
    golden = next(x for x in fixtures['ocr_quality']['goldens'] if x['fixture_id'] == fixture['id'])
    original = ROOT / 'fixtures' / fixture['path']
    assert digest(original) == golden['fixture_sha256']
    image_path = root / 'scan.png'
    shutil.copyfile(original, image_path)
    image = image_path.read_bytes()
    encoded = base64.b64encode(image).decode()

    doc = Document()
    doc.add_paragraph(NATIVE)
    doc.add_picture(str(image_path), width=Inches(6))
    doc.save(root / 'scan.docx')
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), SlideInches(0.5), SlideInches(1), width=SlideInches(8))
    slide.shapes.add_textbox(0, 0, SlideInches(8), SlideInches(1)).text = NATIVE
    slides.save(root / 'scan.pptx')
    book = Workbook()
    book.active['A1'] = NATIVE
    book.active.add_image(Image(str(image_path)), 'A3')
    book.save(root / 'scan.xlsx')
    for family in ['odt', 'ods', 'odp']:
        odf(root / ('scan.' + family), image, family)
    (root / 'scan.html').write_text(f'<html><body><p>{NATIVE}</p><img src="data:image/png;base64,{encoded}" alt="Scan"/></body></html>')
    (root / 'scan.rtf').write_text('{\\rtf1\\ansi ' + NATIVE + '\\par{\\pict\\pngblip ' + image.hex() + '}}')
    notebook = {'nbformat': 4, 'nbformat_minor': 5, 'metadata': {}, 'cells': [
        {'cell_type': 'code', 'id': 'scan', 'metadata': {}, 'source': NATIVE,
         'execution_count': 1, 'outputs': [{'output_type': 'display_data', 'metadata': {},
                                         'data': {'image/png': encoded}}]}]}
    (root / 'scan.ipynb').write_text(json.dumps(notebook))
    archive(root / 'scan.zip', {'scan.docx': (root / 'scan.docx').read_bytes(), 'body.txt': NATIVE})
    archive(root / 'scan.epub', {
        'mimetype': 'application/epub+zip',
        'META-INF/container.xml': '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
        'content.opf': '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">ocr-contract</dc:identifier><dc:title>Scan</dc:title><dc:language>en</dc:language></metadata><manifest><item id="body" href="body.xhtml" media-type="application/xhtml+xml"/><item id="image" href="scan.png" media-type="image/png"/><item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/></manifest><spine toc="ncx"><itemref idref="body"/></spine></package>',
        'body.xhtml': f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Scan</title></head><body><p>{NATIVE}</p><img src="scan.png" alt="Scan"/></body></html>',
        'toc.ncx': '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><head/><docTitle><text>Scan</text></docTitle><navMap><navPoint id="chapter" playOrder="1"><navLabel><text>Scan</text></navLabel><content src="body.xhtml"/></navPoint></navMap></ncx>',
        'scan.png': image})
    sys.path.insert(0, str(ROOT / 'fixtures'))
    spec = importlib.util.spec_from_file_location('repository_fixture_generator', ROOT / 'fixtures/generate.py')
    generator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(generator)
    entries = []
    generator.add_mapi_storage(entries, (), [generator.mapi_unicode(0x0037, 'OCR contract'),
        generator.mapi_long(0x3FFD, 1252), generator.mapi_binary(0x1013,
            f'<html><body><p>{NATIVE}</p><img src="cid:scan@test"/></body></html>'.encode())],
        True, attachments=1)
    generator.add_mapi_storage(entries, ('__attach_version1.0_#00000000',), [
        generator.mapi_long(0x3705, 1), generator.mapi_unicode(0x3707, 'scan.png'),
        generator.mapi_unicode(0x370E, 'image/png'), generator.mapi_unicode(0x3712, 'scan@test'),
        generator.mapi_binary(0x3701, image)], False)
    (root / 'scan.msg').write_bytes(generator.cfb(entries))
    if office:
        for source, extension in [('docx', 'doc'), ('pptx', 'ppt'), ('xlsx', 'xls')]:
            subprocess.run([office, '--headless', '--convert-to', extension, '--outdir', str(root),
                            str(root / ('scan.' + source))], check=True, timeout=120,
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            if not (root / ('scan.' + extension)).is_file():
                raise RuntimeError('independent Office fixture export failed: ' + extension)
    papers, cases = [], []
    for path in sorted(root.glob('scan.*')):
        family = 'image' if path.suffix == '.png' else path.suffix[1:]
        identity = 'ocr-contract-' + family
        papers.append({'id': identity, 'format': family, 'sourcePath': path.name, 'pages': 0,
                       'bytes': path.stat().st_size, 'sha256': digest(path)})
        for mode in ['auto', 'off']:
            case = {'id': identity, 'ocr': mode, 'sourceSha256': digest(path),
                    'requiredText': ([NATIVE] if family != 'image' else []),
                    'forbiddenText': ['\u0001'],
                    'ocrMinimums': {'imageSources': 1}, 'ocrMaximums': {'imagesFailed': 0}}
            if mode == 'auto':
                case['requiredText'].append(golden['ground_truth_nfc'])
                case['exactTextCounts'] = {golden['ground_truth_nfc']: 1}
                case['ocrMinimums'].update(imagesCompleted=1, imagesWithText=1)
                case['ocrMaximums']['imagesSkipped'] = 0
            else:
                case['ocrMinimums']['imagesSkipped'] = 1
                case['ocrMaximums']['imagesAttempted'] = 0
                if family == 'image':
                    case['reviewPages'] = [1]
            cases.append(case)
    (root / 'manifest.json').write_text(json.dumps({'papers': papers}, indent=2) + '\n')
    (root / 'expectations.json').write_text(json.dumps({'cases': cases,
        'imageAuthoritySha256': golden['fixture_sha256'], 'imageOnlyText': golden['ground_truth_nfc']}, indent=2) + '\n')
    print(f'{len(papers)} image-bearing format cases; {len(cases)} OCR mode checks')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--office', help='Optional local LibreOffice executable for DOC/PPT/XLS fixtures')
    args = parser.parse_args()
    build(args.output, args.office)
